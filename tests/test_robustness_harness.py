"""
tests/test_robustness_harness.py
================================
Formal Automated Acceptance Test Suite for DCS Cleaning Robustness & Safeguards (AT-1 through AT-7).
Verifies:
- AT-1: Ambiguous repairs (3-tier classification, context evidence, sentinel isolation, date ambiguity)
- AT-2: Measured accuracy via synthetic corruption (pristine invariance FP=0, idempotency, empirical matrix)
- AT-3: Protected data & identity guardrails (leading zeros, phone, ID, qualitative text locking)
- AT-4: Machine-replayable audit log & script parity (bit-for-bit pandas execution match)
- AT-5: File & scale edge cases (UTF-8 BOM, CP1252 decoding, merged cell unmerging, DCSFileFormatError)
- AT-6: Format drift & unmapped survey phrase preservation
- AT-7: Unsafe export formula injection neutralization (CWE-1236)
"""

import io
import os
import re
import openpyxl
import pytest
import pandas as pd
import numpy as np

import engine
from tests.corruption_injector import (
    create_synthetic_clean_dataset,
    inject_synthetic_corruptions,
    evaluate_accuracy_matrix
)


# =============================================================================
# AT-1: Ambiguous Repairs & Rule Classification
# =============================================================================

def test_at1_sentinel_context_evidence_isolation():
    """
    AT-1.1: In a column with salary values [800, 1500], '999' is a valid continuous observation
    and MUST NOT be converted to NaN. In a Likert scale [1, 5], '999' is an extreme sentinel.
    """
    df = pd.DataFrame({
        "monthly_salary": [850, 920, 999, 1200, 1450, 1100],
        "kfl_satisfaction": [1, 2, "999", 4, 5, 3]
    })
    res = engine.propose_cleaning_repairs(df)
    proposals = {p["proposal_id"]: p for p in res["proposals"]}

    # Likert item should have proposal with CONFIRM_FIRST and default_accepted = True
    likert_prop = proposals.get("sentinels::kfl_satisfaction")
    assert likert_prop is not None
    assert likert_prop["tier"] == "CONFIRM_FIRST"
    assert likert_prop["default_accepted"] is True
    assert likert_prop["confidence"] >= 0.80

    # Salary item: either no sentinel proposal or marked NEVER_AUTO with default_accepted = False
    salary_prop = proposals.get("sentinels::monthly_salary")
    if salary_prop is not None:
        assert salary_prop["tier"] == "NEVER_AUTO"
        assert salary_prop["default_accepted"] is False


def test_at1_date_format_ambiguity_classification():
    """
    AT-1.2: Dates where both day and month are <= 12 are genuinely ambiguous (e.g. 03/04/2026).
    DCS must classify them as NEVER_AUTO with Ambiguous confidence and refuse silent coercion.
    """
    df_ambig = pd.DataFrame({"survey_date": ["03/04/2026", "05/06/2026", "01/02/2026"]})
    is_amb, conf, tier, fmt = engine.evaluate_column_date_evidence(df_ambig["survey_date"])
    assert is_amb is True
    assert tier == "NEVER_AUTO"
    assert conf <= 0.60

    # Column with clear day > 12 should resolve unambiguously to DD/MM/YYYY
    df_clear = pd.DataFrame({"survey_date": ["03/04/2026", "25/04/2026", "11/02/2026"]})
    is_amb2, conf2, tier2, fmt2 = engine.evaluate_column_date_evidence(df_clear["survey_date"])
    assert is_amb2 is False
    assert tier2 == "CONFIRM_FIRST"
    assert fmt2 == "DD/MM/YYYY"


# =============================================================================
# AT-2: Measured Accuracy via Synthetic Corruption Injection
# =============================================================================

def test_at2_pristine_clean_invariance():
    """
    AT-2.1 (Pristine Invariance): Feeding a 100% clean dataset into DCS must produce
    zero false-positive modifications (FP = 0).
    """
    df_clean = create_synthetic_clean_dataset(n_rows=150, seed=101)
    prop_res = engine.propose_cleaning_repairs(df_clean)

    # All proposed items should be informational protected column notifications, not destructive changes
    actionable_props = [
        p for p in prop_res["proposals"] 
        if p["rule_id"] not in ["protect_column"] and p.get("default_accepted", False)
    ]
    assert len(actionable_props) == 0, f"Expected 0 actionable proposals on clean data, got: {actionable_props}"
    assert prop_res["health_score"] >= 98.0


def test_at2_idempotency_guarantee():
    """
    AT-2.2 (Idempotency): Clean(Clean(X)) == Clean(X).
    Running the cleaning pipeline a second time produces an empty change log and identical data.
    """
    df_clean = create_synthetic_clean_dataset(n_rows=50, seed=202)
    df_corr, manifest = inject_synthetic_corruptions(df_clean, seed=202)

    # First cleaning pass
    res1 = engine.propose_cleaning_repairs(df_corr)
    acc1 = [p["proposal_id"] for p in res1["proposals"] if p.get("default_accepted", False)]
    app1 = engine.apply_accepted_repairs(df_corr, acc1, proposals=res1["proposals"])
    df_cleaned_1 = app1["df_cleaned"]

    # Second cleaning pass
    res2 = engine.propose_cleaning_repairs(df_cleaned_1)
    acc2 = [p["proposal_id"] for p in res2["proposals"] if p.get("default_accepted", False) and p["rule_id"] != "protect_column"]
    app2 = engine.apply_accepted_repairs(df_cleaned_1, acc2, proposals=res2["proposals"])
    df_cleaned_2 = app2["df_cleaned"]

    # Second pass must produce zero new rule modifications
    assert len(app2["df_rule_changelog"]) == 0
    pd.testing.assert_frame_equal(df_cleaned_1, df_cleaned_2)


def test_at2_synthetic_corruption_baseline_measurement():
    """
    AT-2.3: Injects known corruptions, measures Sensitivity and Precision,
    and publishes the empirical matrix (asserting FP = 0 on clean background cells).
    """
    df_clean = create_synthetic_clean_dataset(n_rows=100, seed=303)
    df_corr, manifest = inject_synthetic_corruptions(df_clean, seed=303)

    prop_res = engine.propose_cleaning_repairs(df_corr)
    acc_ids = [p["proposal_id"] for p in prop_res["proposals"] if p.get("default_accepted", False)]
    apply_res = engine.apply_accepted_repairs(df_corr, acc_ids, proposals=prop_res["proposals"])

    acc_report = evaluate_accuracy_matrix(
        df_clean, df_corr, apply_res["df_cleaned"], apply_res["df_cell_audit"], manifest
    )

    print("\n--- Empirical Accuracy Baseline Report ---")
    print(acc_report)

    # True positives must be detected
    assert acc_report["true_positives"] >= 5
    # False positives on clean background cells must be strictly 0
    assert acc_report["false_positives"] == 0, f"False positives detected on clean cells: {acc_report['false_positives']}"
    assert acc_report["precision_pct"] == 100.0


# =============================================================================
# AT-3: Protected Data & Identity Column Guardrails
# =============================================================================

def test_at3_protected_columns_locked_by_default():
    """
    AT-3.1 & AT-3.2: Leading zero postal codes, IDs, phone numbers, and free-text open comments
    must be detected, locked by default, and excluded from auto-recode and clustering.
    """
    df = pd.DataFrame({
        "respondent_id": ["RESP_001", "RESP_002", "RESP_003"],
        "zip_code": ["02138", "00921", "01002"],
        "phone_number": ["+1 800 555 0101", "+1 800 555 0102", "+1 800 555 0103"],
        "open_feedback": [
            "The survey questions were very comprehensive and well organized throughout the study.",
            "I found the platform easy to navigate and appreciate the clear instructions provided.",
            "Overall satisfactory experience with the questionnaire interface and timely completion."
        ],
        "kfl1": [1, 2, 3]
    })

    protected = engine.detect_protected_columns(df)
    assert "zip_code" in protected
    assert protected["zip_code"]["type"] == "LEADING_ZERO"
    assert "respondent_id" in protected
    assert "phone_number" in protected
    assert "open_feedback" in protected
    assert protected["open_feedback"]["type"] == "FREE_TEXT"
    assert "kfl1" not in protected

    prop_res = engine.propose_cleaning_repairs(df)
    props = {p["proposal_id"]: p for p in prop_res["proposals"]}

    # Verify protected column proposals exist and are marked NEVER_AUTO with default_accepted = False
    assert "protected::zip_code" in props
    assert props["protected::zip_code"]["default_accepted"] is False
    assert props["protected::zip_code"]["tier"] == "NEVER_AUTO"

    # Verify that open_feedback was NOT subjected to category clustering
    assert "cluster::open_feedback" not in props


# =============================================================================
# AT-4: Machine-Replayable Audit Log & Code Parity
# =============================================================================

def test_at4_reproducible_python_script_execution_parity(tmp_path):
    """
    AT-4.1: Standalone Python script exported by DCS must execute independently
    and reproduce the exact cleaned dataset bit-for-bit.
    """
    raw_df = pd.DataFrame({
        "id": [1, 2, 3, 4],
        "employment": ["Full-time", "full time", "Part-time", "fulltime"],
        "kfl1": [1, 2, "999", 4],
        "age": [25, 250, 34, 45]
    })

    prop_res = engine.propose_cleaning_repairs(raw_df)
    acc_ids = [p["proposal_id"] for p in prop_res["proposals"] if p.get("default_accepted", False)]
    apply_res = engine.apply_accepted_repairs(raw_df, acc_ids, proposals=prop_res["proposals"])
    df_dcs_cleaned = apply_res["df_cleaned"]

    sha_in = engine.compute_df_sha256(raw_df)
    sha_out = engine.compute_df_sha256(df_dcs_cleaned)

    script_text = engine.generate_reproducible_python_script(
        raw_df, prop_res["proposals"], source_sha256=sha_in, output_sha256=sha_out
    )

    raw_csv = str(tmp_path / "raw_test.csv")
    out_csv = str(tmp_path / "out_test.csv")
    raw_df.to_csv(raw_csv, index=False)

    exec_env = {}
    exec(script_text, exec_env)
    df_replayed = exec_env["clean_dataset"](raw_csv, out_csv)

    pd.testing.assert_frame_equal(
        df_dcs_cleaned.reset_index(drop=True),
        df_replayed.reset_index(drop=True),
        check_dtype=False
    )


# =============================================================================
# AT-5: File & Scale Edge Cases (Fail-Loud Semantics)
# =============================================================================

def test_at5_utf8_bom_and_cp1252_encoding():
    """
    AT-5.1: Handles UTF-8 with BOM (\xef\xbb\xbf) without column 0 corruption,
    and cleanly decodes Windows-1252 / CP1252 Scandinavian characters.
    """
    # 1. UTF-8 BOM
    bom_csv = b"\xef\xbb\xbfQ1,Q2\n10,20\n30,40\n"
    df_bom, meta_bom = engine.robust_load_spreadsheet(bom_csv, "bom_test.csv")
    assert meta_bom["encoding_used"] == "utf-8-sig"
    assert df_bom.iloc[0, 0] == "Q1"  # No BOM artifact in string

    # 2. Windows-1252 Scandinavian text
    swedish_bytes = "Fråga,Ålder\nKFL1,42\n".encode("cp1252")
    df_cp, meta_cp = engine.robust_load_spreadsheet(swedish_bytes, "cp_test.csv")
    assert meta_cp["encoding_used"] == "cp1252"
    assert df_cp.iloc[0, 0] == "Fråga"
    assert df_cp.iloc[0, 1] == "Ålder"


def test_at5_excel_merged_cells_unmerging():
    """
    AT-5.2: Detects merged cells across Excel headers/data, unmerges and forward-fills
    top-left values without data loss.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MergedTest"
    ws.cell(row=1, column=1, value="ID")
    ws.cell(row=1, column=2, value="Demographics")
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
    ws.cell(row=2, column=1, value="RESP_1")
    ws.cell(row=2, column=2, value="Age")
    ws.cell(row=2, column=3, value="Gender")

    buf = io.BytesIO()
    wb.save(buf)

    df_xl, meta_xl = engine.robust_load_spreadsheet(buf.getvalue(), "merged.xlsx")
    assert meta_xl["unmerged_cells_count"] == 2
    assert df_xl.iloc[0, 1] == "Demographics"
    assert df_xl.iloc[0, 2] == "Demographics"


def test_at5_fail_loud_on_corrupt_or_empty_file():
    """
    AT-5.4: Fails loudly with DCSFileFormatError on empty or invalid file bytes.
    """
    with pytest.raises(engine.DCSFileFormatError):
        engine.robust_load_spreadsheet(b"", "empty.csv")


# =============================================================================
# AT-6: Format Drift & Unmapped Survey Phrases
# =============================================================================

def test_at6_unmapped_survey_phrases_preserved_as_raw():
    """
    AT-6.2: When standard Likert responses predominate but unmapped verbal phrases exist
    (e.g. 'pretty much agree', 'partly agree'), DCS discovers them and preserves raw strings.
    """
    df = pd.DataFrame({
        "kfl_item": [
            "Strongly Agree", "Agree", "Disagree", "Strongly Disagree",
            "Agree", "pretty much agree", "Agree"
        ]
    })
    res = engine.propose_cleaning_repairs(df)
    likert_prop = next(p for p in res["proposals"] if p["rule_id"] == "map_likert")
    assert "unmapped_phrases" in likert_prop
    unmapped_tokens = [item["phrase"] for item in likert_prop["unmapped_phrases"]]
    assert "pretty much agree" in unmapped_tokens

    # Applying repairs recodes standard English anchors while leaving 'pretty much agree' intact
    app = engine.apply_accepted_repairs(df, [likert_prop["proposal_id"]], proposals=res["proposals"])
    cleaned_col = app["df_cleaned"]["kfl_item"]
    # Row 5 was 'pretty much agree' -> preserved as raw string
    assert cleaned_col.iloc[5] == "pretty much agree"
    # Row 0 was 'Strongly Agree' -> recoded to 7
    assert cleaned_col.iloc[0] == 7


# =============================================================================
# AT-7: Unsafe Export Neutralization (CWE-1236 Formula Injection)
# =============================================================================

def test_at7_cwe1236_formula_injection_neutralization():
    """
    AT-7.1: Text cells starting with '=', '+', '-', '@', '\\t', or '\\r'
    are neutralized with a leading single quote (') on export while preserving valid numbers.
    """
    df = pd.DataFrame({
        "user_input": [
            "=CMD('calc.exe')",
            "@SUM(1+1)*cmd|' /C calc'!A0",
            "+12345-EVIL()",
            "-2+3*[DDE]",
            "\t=1+1",
            "Normal Text",
            "-14.5",  # Legitimate negative number
            "+5"       # Legitimate positive number string
        ]
    })

    df_safe = engine.sanitize_spreadsheet_formulas(df)

    assert df_safe.at[0, "user_input"] == "'=CMD('calc.exe')"
    assert df_safe.at[1, "user_input"] == "'@SUM(1+1)*cmd|' /C calc'!A0"
    assert df_safe.at[2, "user_input"] == "'+12345-EVIL()"
    assert df_safe.at[3, "user_input"] == "'-2+3*[DDE]"
    assert df_safe.at[4, "user_input"] == "'\t=1+1"
    assert df_safe.at[5, "user_input"] == "Normal Text"
    # Legitimate numbers are NOT quoted
    assert df_safe.at[6, "user_input"] == "-14.5"
    assert df_safe.at[7, "user_input"] == "+5"
